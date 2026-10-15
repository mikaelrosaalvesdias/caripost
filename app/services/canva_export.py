"""Exporta uma arte EM CAMADAS (layered_composer) para o Canva como design EDITÁVEL e traz de volta.

Fluxo: layers → PPTX (cada camada = objeto: foto, bloco, texto com fonte Poppins, logo) → URL pública →
Composio CANVA_POST_URL_IMPORTS → design editável (link de edição) → depois CANVA_POST_EXPORTS (PNG) → baixa.
A conta Canva usada é a conectada na Composio (setting canva.connected_account_id / canva.user_id; se vazio,
pega a 1ª conta Canva ACTIVE)."""
from __future__ import annotations

import asyncio
import io
from typing import Any, Dict, List, Optional, Tuple

PX_PER_IN = 288          # 2160 px = 7,5 pol (o Canva mantém a proporção; 1 px do canvas = 1/4 pt)
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def layers_to_pptx(slides: List[Tuple[int, int, List[Dict[str, Any]]]]) -> bytes:
    """slides = [(W, H, layers), ...] → PPTX com 1 slide por tela (carrossel = várias páginas no Canva)."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR
    from pptx.util import Emu, Pt

    W, H = slides[0][0], slides[0][1]
    emu = lambda px: Emu(int(px * 914400 / PX_PER_IN))
    prs = Presentation()
    prs.slide_width, prs.slide_height = emu(W), emu(H)
    blank = prs.slide_layouts[6]
    for (_w, _h, layers) in slides:
        sl = prs.slides.add_slide(blank)
        for ly in layers:
            k = ly.get("kind")
            x, y, w, h = emu(ly["x"]), emu(ly["y"]), emu(max(1, ly["w"])), emu(max(1, ly["h"]))
            if k == "rect":
                shp = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if ly.get("radius") else MSO_SHAPE.RECTANGLE, x, y, w, h)
                if ly.get("radius"):
                    try:
                        shp.adjustments[0] = min(0.5, ly["radius"] / max(1, min(ly["w"], ly["h"])))
                    except Exception:
                        pass
                shp.fill.solid(); shp.fill.fore_color.rgb = RGBColor.from_string(ly["fill"].lstrip("#"))
                shp.line.fill.background(); shp.shadow.inherit = False
                shp.name = ly.get("name") or "bloco"
            elif k == "image":
                pic = sl.shapes.add_picture(io.BytesIO(ly["data"]), x, y, w, h)
                pic.name = ly.get("name") or "imagem"
            elif k == "text":
                tb = sl.shapes.add_textbox(x, y, emu(ly["w"] + ly["size"] * 0.6), h)
                tb.name = ly.get("name") or "texto"
                tf = tb.text_frame
                tf.word_wrap = True
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.vertical_anchor = MSO_ANCHOR.TOP
                for i, line in enumerate((ly.get("text") or "").split("\n")):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p.line_spacing = float(ly.get("line_height") or 1.12)
                    r = p.add_run(); r.text = line
                    f = r.font
                    f.name = ly.get("font") or "Poppins"
                    f.size = Pt(ly["size"] * 72 / PX_PER_IN)
                    f.bold = ly.get("weight") in ("semibold", "bold")
                    f.color.rgb = RGBColor.from_string(ly["color"].lstrip("#"))
    buf = io.BytesIO(); prs.save(buf)
    return buf.getvalue()


async def _provider():
    from ..social.composio_provider import ComposioConfig, ComposioSocialProvider
    try:
        return ComposioSocialProvider(ComposioConfig.from_env())
    except TypeError:
        return ComposioSocialProvider()


async def _canva_account(p) -> Tuple[str, str]:
    from .ai import get_setting
    acc, uid = (get_setting("canva.connected_account_id") or "").strip(), (get_setting("canva.user_id") or "").strip()
    if acc and uid:
        return acc, uid
    r = await p._request("GET", "/connected_accounts?toolkit_slugs=canva")
    for it in r.get("items") or []:
        if (it.get("status") or "").upper() == "ACTIVE":
            return it["id"], it.get("user_id") or ""
    raise ValueError("Canva não conectado")


async def _exec(p, slug: str, args: Dict[str, Any]) -> Dict[str, Any]:
    acc, uid = await _canva_account(p)
    r = await p._request("POST", f"/tools/execute/{slug}", json={"connected_account_id": acc, "user_id": uid,
                                                                 "arguments": args, "allow_tracing": False})
    if r.get("successful") is False:
        raise ValueError(f"Canva ({slug}): {str(r.get('error') or r)[:300]}")
    return r.get("data") or r


def _dig(d: Any, *keys) -> Any:
    for k in keys:
        if isinstance(d, dict):
            d = d.get(k)
        else:
            return None
    return d


async def import_to_canva(pptx_url: str, title: str) -> Dict[str, Any]:
    """Importa o PPTX (URL pública) → {design_id, edit_url, view_url}."""
    p = await _provider()
    d = await _exec(p, "CANVA_POST_URL_IMPORTS", {"url": pptx_url, "title": title[:250], "mime_type": PPTX_MIME})
    job = _dig(d, "job") or d
    jid = job.get("id")
    for _ in range(60):
        st = job.get("status")
        if st == "success":
            ds = (_dig(job, "result", "designs") or [{}])[0]
            return {"design_id": ds.get("id"), "edit_url": _dig(ds, "urls", "edit_url"),
                    "view_url": _dig(ds, "urls", "view_url"), "thumbnail": _dig(ds, "thumbnail", "url")}
        if st == "failed":
            raise ValueError("Canva recusou o arquivo: " + str(_dig(job, "error", "message") or job.get("error"))[:240])
        await asyncio.sleep(3)
        job = _dig(await _exec(p, "CANVA_GET_URL_IMPORTS_JOBID", {"jobId": jid}), "job") or {}
    raise ValueError("O Canva demorou demais para importar.")


async def export_from_canva(design_id: str) -> List[bytes]:
    """Exporta o design do Canva em PNG (todas as páginas, tamanho original) → lista de imagens."""
    import httpx
    p = await _provider()
    d = await _exec(p, "CANVA_POST_EXPORTS", {"design_id": design_id, "format": {"type": "png", "lossless": True}})
    job = _dig(d, "job") or d
    jid = job.get("id")
    for _ in range(60):
        st = job.get("status")
        if st == "success":
            out = []
            async with httpx.AsyncClient(timeout=120) as c:
                for u in job.get("urls") or []:
                    out.append((await c.get(u)).content)
            return out
        if st == "failed":
            raise ValueError("Canva não exportou: " + str(_dig(job, "error", "message") or job.get("error"))[:240])
        await asyncio.sleep(3)
        job = _dig(await _exec(p, "CANVA_GET_DESIGN_EXPORT_JOB_RESULT", {"exportId": jid}), "job") or {}
    raise ValueError("O Canva demorou demais para exportar.")


async def canva_status() -> Dict[str, Any]:
    p = await _provider()
    r = await p._request("GET", "/connected_accounts?toolkit_slugs=canva")
    items = r.get("items") or []
    act = [i for i in items if (i.get("status") or "").upper() == "ACTIVE"]
    return {"connected": bool(act), "accounts": [{"id": i.get("id"), "status": i.get("status")} for i in items]}


async def canva_connect_link(callback_url: str = "") -> str:
    """Link de autorização OAuth do Canva (Composio). Depois de autorizar, a conta fica ACTIVE e é usada sozinha."""
    from .ai import set_setting
    p = await _provider()
    r = await p._request("GET", "/auth_configs?toolkit_slug=canva")
    cfgs = [i for i in (r.get("items") or []) if (i.get("status") or "").upper() == "ENABLED"]
    if not cfgs:
        raise ValueError("Não há configuração de login do Canva na Composio.")
    body: Dict[str, Any] = {"auth_config_id": cfgs[0]["id"], "user_id": "caripost-canva"}
    if callback_url:
        body["callback_url"] = callback_url
    d = await p._request("POST", "/connected_accounts/link", json=body)
    for k in ("canva.connected_account_id", "canva.user_id"):
        try:
            set_setting(k, "")          # esquece a conta antiga; a nova (ACTIVE) é descoberta sozinha
        except Exception:
            pass
    if not d.get("redirect_url"):
        raise ValueError("A Composio não devolveu o link de conexão.")
    return d["redirect_url"]
